"""Load and normalize the practice vocabulary from the source workbooks."""

from __future__ import annotations

from collections import OrderedDict
from pathlib import Path
import re
import unicodedata

from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parents[2]
DOC_DIR = PROJECT_ROOT / "doc"


def normalize_answer(value: str) -> str:
    """Normalize user-entered Japanese for lenient but exact matching."""

    value = unicodedata.normalize("NFKC", value or "")
    return re.sub(r"\s+", "", value).strip().lower()


def _is_unverified_accent(value: str) -> bool:
    return not value or "待核" in value


def _rows(path: Path) -> list[dict[str, str]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.worksheets[0]
    values = sheet.iter_rows(values_only=True)
    headers = [str(value).strip() for value in next(values)]
    records: list[dict[str, str]] = []
    for row in values:
        record = {
            header: "" if value is None else str(value).strip()
            for header, value in zip(headers, row, strict=False)
        }
        if any(record.values()):
            records.append(record)
    workbook.close()
    return records


def _lesson(value: str) -> int | str:
    try:
        return int(float(value))
    except ValueError:
        return value


def _merge_unique(target: list, values: list) -> None:
    for value in values:
        if value not in target:
            target.append(value)


def _transitivity(notes: list[str]) -> str:
    joined = "；".join(notes)
    if "自/他" in joined or "他/自" in joined or "视搭配" in joined:
        return "自动词／他动词（视搭配）"
    markers = {note[:1] for note in notes if note[:1] in {"自", "他"}}
    if markers == {"自", "他"}:
        return "自动词／他动词"
    if "自" in markers:
        return "自动词"
    if "他" in markers:
        return "他动词"
    return "未标注"


def _adjective_answers(original: str, kana: str, adjective_type: str) -> list[str]:
    candidates = [original, kana]
    if adjective_type == "na形容词":
        base_original = re.sub(r"[（(]な[）)]$", "", original)
        candidates.extend([base_original, f"{base_original}な", f"{kana}な"])
    return list(dict.fromkeys(normalize_answer(value) for value in candidates))


def _kana_masu_form(kana: str, verb_type: str) -> str:
    if verb_type.startswith("二类") and kana.endswith("る"):
        return f"{kana[:-1]}ます"
    if "サ变" in verb_type and kana.endswith("する"):
        return f"{kana[:-2]}します"
    if "カ变" in verb_type and kana.endswith("くる"):
        return f"{kana[:-2]}きます"
    godan_endings = {
        "う": "います",
        "く": "きます",
        "ぐ": "ぎます",
        "す": "します",
        "つ": "ちます",
        "ぬ": "にます",
        "ぶ": "びます",
        "む": "みます",
        "る": "ります",
    }
    ending = godan_endings.get(kana[-1:])
    return f"{kana[:-1]}{ending}" if ending else ""


def _load_verbs() -> list[dict]:
    records: OrderedDict[tuple[str, str], dict] = OrderedDict()
    path = DOC_DIR / "新标日初级上册动词表.xlsx"
    for row in _rows(path):
        original = row["原型（汉字+假名）"]
        kana = row["假名"]
        # The source has one malformed row whose "original" is すみます and
        # whose meaning mixes 住む with 済む. It cannot produce one valid answer.
        if original == "すみます" and row["masu型"] == "すみます":
            continue
        key = (original, kana)
        item = records.setdefault(
            key,
            {
                "kind": "verb",
                "original": original,
                "kana": kana,
                "meanings": [],
                "lessons": [],
                "accent": None,
                "verb_type": row["几型动词"],
                "masu_form": row["masu型"],
                "te_form": row["te型"],
                "notes": [],
            },
        )
        _merge_unique(item["meanings"], [row["中文意思"]])
        _merge_unique(item["lessons"], [_lesson(row["课次"])])
        _merge_unique(item["notes"], [row["备注"]])
        if not _is_unverified_accent(row["音调"]):
            item["accent"] = row["音调"]

    result = []
    for index, item in enumerate(records.values(), start=1):
        item["id"] = f"verb-{index}"
        item["meaning"] = "；".join(item.pop("meanings"))
        item["note"] = "；".join(item.pop("notes"))
        item["transitivity"] = _transitivity(item["note"].split("；"))
        accepted_answers = (
            item["original"],
            item["kana"],
            item["masu_form"],
            _kana_masu_form(item["kana"], item["verb_type"]),
        )
        item["answers"] = list(
            dict.fromkeys(normalize_answer(value) for value in accepted_answers if value)
        )
        result.append(item)
    return result


def _load_adjectives() -> list[dict]:
    records: OrderedDict[tuple[str, str], dict] = OrderedDict()
    path = DOC_DIR / "新标日初级上册形容词表.xlsx"
    for row in _rows(path):
        original = row["原型（汉字＋假名）"]
        kana = row["假名"]
        key = (original, kana)
        adjective_type = "i形容词" if row["词性"].startswith("一类") else "na形容词"
        item = records.setdefault(
            key,
            {
                "kind": "adjective",
                "original": original,
                "kana": kana,
                "meanings": [],
                "lessons": [],
                "accent": None,
                "adjective_type": adjective_type,
                "conjunctive_form": row["连用形（く／に）"],
                "past_form": row["过去式（かった／だった）"],
                "notes": [],
            },
        )
        _merge_unique(item["meanings"], [row["中文意思"]])
        _merge_unique(item["lessons"], [_lesson(row["课次"])])
        _merge_unique(item["notes"], [row["出处／备注"]])
        if not _is_unverified_accent(row["音调"]):
            item["accent"] = row["音调"]

    result = []
    for index, item in enumerate(records.values(), start=1):
        item["id"] = f"adjective-{index}"
        item["meaning"] = "；".join(item.pop("meanings"))
        item["note"] = "；".join(item.pop("notes"))
        item["answers"] = _adjective_answers(
            item["original"], item["kana"], item["adjective_type"]
        )
        result.append(item)
    return result


def load_words() -> list[dict]:
    """Return all deduplicated practice words."""

    return _load_verbs() + _load_adjectives()
